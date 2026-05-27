import io

import openpyxl
from docx import Document
from pypdf import PdfReader

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SUPPORTED_MIME_TYPES = {PDF_MIME, DOCX_MIME, XLSX_MIME}
MAX_UPLOAD_BYTES = 10 * 1024 * 1024


def extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


def extract_text_from_xlsx(file_bytes: bytes) -> str:
    workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            parts.extend(str(cell) for cell in row if cell is not None)
    return " ".join(parts)


def extract_text(file_bytes: bytes, mime_type: str) -> str:
    if mime_type == PDF_MIME:
        return extract_text_from_pdf(file_bytes)
    if mime_type == DOCX_MIME:
        return extract_text_from_docx(file_bytes)
    if mime_type == XLSX_MIME:
        return extract_text_from_xlsx(file_bytes)
    raise ValueError("Unsupported file type")


def is_supported_file(filename: str, mime_type: str, allow_xlsx: bool = False) -> bool:
    extension = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    allowed = {PDF_MIME: "pdf", DOCX_MIME: "docx"}
    if allow_xlsx:
        allowed[XLSX_MIME] = "xlsx"
    return allowed.get(mime_type) == extension

